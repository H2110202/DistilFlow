import json
import os
from pathlib import Path
from typing import Optional
from backend.tools.registry import registry
from backend.config import get_settings


@registry.register(name="scan_folder", category="file")
def scan_folder(folder_path: str, recursive: bool = True, extensions: str = "") -> str:
    settings = get_settings()
    target = Path(folder_path)
    if not target.exists():
        target = Path(settings.data_dir) / folder_path
    if not target.exists():
        return json.dumps({"error": f"文件夹不存在: {folder_path}"}, ensure_ascii=False)
    if not target.is_dir():
        return json.dumps({"error": f"不是文件夹: {folder_path}"}, ensure_ascii=False)

    ext_filter = [e.strip().lower().lstrip(".") for e in extensions.split(",") if e.strip()] if extensions else []
    files = []
    glob_fn = target.rglob if recursive else target.glob

    for f in glob_fn("*"):
        if not f.is_file():
            continue
        ext = f.suffix.lower().lstrip(".")
        if ext_filter and ext not in ext_filter:
            continue
        rel_path = str(f.relative_to(target))
        file_type = _detect_file_type(f)
        files.append({
            "name": f.name,
            "path": str(f),
            "relative_path": rel_path,
            "extension": ext,
            "file_type": file_type,
            "size_bytes": f.stat().st_size,
        })

    type_counts = {}
    for f in files:
        t = f["file_type"]
        type_counts[t] = type_counts.get(t, 0) + 1

    return json.dumps({
        "folder": str(target),
        "total_files": len(files),
        "type_counts": type_counts,
        "files": files,
    }, ensure_ascii=False, default=str)


def _detect_file_type(f: Path) -> str:
    ext = f.suffix.lower()
    excel_exts = {".xlsx", ".xls", ".xlsm", ".xlsb", ".csv"}
    pdf_exts = {".pdf"}
    image_exts = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"}
    word_exts = {".doc", ".docx"}
    if ext in excel_exts:
        return "excel"
    if ext in pdf_exts:
        return "pdf"
    if ext in image_exts:
        return "image"
    if ext in word_exts:
        return "word"
    return "other"


@registry.register(name="analyze_template", category="file")
def analyze_template(file_path: str, sheet_name: str = "") -> str:
    import openpyxl
    from openpyxl.utils import get_column_letter

    path = Path(file_path)
    if not path.exists():
        settings = get_settings()
        path = Path(settings.data_dir) / file_path
    if not path.exists():
        return json.dumps({"error": f"文件不存在: {file_path}"}, ensure_ascii=False)

    ext = path.suffix.lower()
    if ext not in {".xlsx", ".xls", ".xlsm", ".csv"}:
        return json.dumps({"error": f"模板分析目前仅支持 Excel 文件，收到: {ext}"}, ensure_ascii=False)

    try:
        wb = openpyxl.load_workbook(str(path))
        sheets_to_analyze = [sheet_name] if sheet_name else wb.sheetnames
        all_fields = []

        for sname in sheets_to_analyze:
            if sname not in wb.sheetnames:
                continue
            ws = wb[sname]

            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None or str(cell.value).strip() == "":
                        has_adjacent_label = False
                        label = ""
                        if cell.column > 1:
                            left = ws.cell(row=cell.row, column=cell.column - 1)
                            if left.value and str(left.value).strip():
                                has_adjacent_label = True
                                label = str(left.value).strip()
                        if not has_adjacent_label and cell.row > 1:
                            above = ws.cell(row=cell.row - 1, column=cell.column)
                            if above.value and str(above.value).strip():
                                has_adjacent_label = True
                                label = str(above.value).strip()
                        if has_adjacent_label:
                            all_fields.append({
                                "sheet": sname,
                                "cell": f"{get_column_letter(cell.column)}{cell.row}",
                                "row": cell.row,
                                "column": cell.column,
                                "label": label,
                                "current_value": "",
                                "type_hint": "text",
                            })
                    elif str(cell.value).strip():
                        val = str(cell.value).strip()
                        is_header_like = (
                            cell.row <= 5
                            or (cell.column > 1 and ws.cell(row=cell.row, column=cell.column - 1).value is None)
                            or len(val) < 20
                        )
                        if is_header_like and not val.replace(".", "").replace("-", "").replace(",", "").isdigit():
                            right_cell = ws.cell(row=cell.row, column=cell.column + 1)
                            below_cell = ws.cell(row=cell.row + 1, column=cell.column) if cell.row < ws.max_row else None
                            if (right_cell.value is None or str(right_cell.value).strip() == "") and right_cell.column <= ws.max_column:
                                all_fields.append({
                                    "sheet": sname,
                                    "cell": f"{get_column_letter(cell.column + 1)}{cell.row}",
                                    "row": cell.row,
                                    "column": cell.column + 1,
                                    "label": val,
                                    "current_value": "",
                                    "type_hint": "text",
                                })

            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and str(cell.value).strip():
                        val = str(cell.value).strip()
                        for field in all_fields:
                            if field["current_value"] == "" and field["row"] == cell.row and field["column"] == cell.column:
                                field["current_value"] = val
                                break

        wb.close()

        return json.dumps({
            "file": str(path),
            "sheets_analyzed": sheets_to_analyze,
            "total_fields": len(all_fields),
            "fields": all_fields,
        }, ensure_ascii=False, default=str)

    except Exception as e:
        return json.dumps({"error": f"模板分析失败: {str(e)}"}, ensure_ascii=False)


@registry.register(name="match_fields", category="intelligence")
def match_fields(source_data: str, template_fields: str) -> str:
    from backend.llm_client import chat

    try:
        src = json.loads(source_data) if isinstance(source_data, str) else source_data
        fields = json.loads(template_fields) if isinstance(template_fields, str) else template_fields
    except:
        return json.dumps({"error": "输入数据格式错误，需要JSON"}, ensure_ascii=False)

    field_list = fields if isinstance(fields, list) else fields.get("fields", [])
    if not field_list:
        return json.dumps({"error": "模板字段为空"}, ensure_ascii=False)

    src_summary = json.dumps(src, ensure_ascii=False, default=str)
    if len(src_summary) > 8000:
        src_summary = src_summary[:8000] + "...(截断)"

    field_descriptions = "\n".join(
        f"  - 字段{i+1}: 标签「{f.get('label', '')}」 位置[{f.get('sheet', '')}!{f.get('cell', '')}]"
        for i, f in enumerate(field_list)
    )

    prompt = f"""你是数据匹配专家。请将源数据中的信息匹配到模板字段中。

源数据摘要:
{src_summary}

模板字段:
{field_descriptions}

请为每个字段找到最匹配的源数据值。如果某个字段在源数据中找不到对应值，填 null。

输出JSON数组，每个元素包含:
- label: 模板字段标签
- cell: 单元格位置
- sheet: 工作表名
- value: 匹配到的值（找不到填 null）
- confidence: 匹配置信度 high/medium/low
- source: 值来自源数据的哪个位置

只输出JSON数组，不要其他文字。"""

    result = chat([{"role": "user", "content": prompt}], temperature=0.1, max_tokens=3000)
    if not result:
        return json.dumps({"error": "LLM匹配失败"}, ensure_ascii=False)

    result = result.strip()
    if result.startswith("```"):
        result = result.split("\n", 1)[-1].rsplit("```", 1)[0].strip()

    try:
        matched = json.loads(result)
    except:
        return json.dumps({"error": "LLM返回格式错误", "raw": result[:500]}, ensure_ascii=False)

    if isinstance(matched, dict) and "fields" in matched:
        matched = matched["fields"]
    elif isinstance(matched, dict):
        matched = [matched]

    return json.dumps({
        "total_fields": len(field_list),
        "matched_count": sum(1 for m in matched if m.get("value") is not None),
        "matches": matched,
    }, ensure_ascii=False, default=str)


@registry.register(name="fill_template", category="file")
def fill_template(template_path: str, fill_data: str, output_filename: str = "") -> str:
    import openpyxl
    from openpyxl.utils import get_column_letter
    import shutil

    path = Path(template_path)
    if not path.exists():
        settings = get_settings()
        path = Path(settings.data_dir) / template_path
    if not path.exists():
        return json.dumps({"error": f"模板文件不存在: {template_path}"}, ensure_ascii=False)

    try:
        data = json.loads(fill_data) if isinstance(fill_data, str) else fill_data
    except:
        return json.dumps({"error": "填充数据格式错误"}, ensure_ascii=False)

    settings = get_settings()
    output_dir = Path(settings.data_dir) / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    out_name = output_filename or f"filled_{path.stem}_{Path(template_path).stat().st_mtime:.0f}.xlsx"
    out_path = output_dir / out_name

    shutil.copy2(str(path), str(out_path))

    try:
        wb = openpyxl.load_workbook(str(out_path))
        filled_count = 0
        failed_cells = []

        if isinstance(data, list):
            for item in data:
                sheet_name = item.get("sheet", "")
                cell_ref = item.get("cell", "")
                value = item.get("value")

                if value is None:
                    continue

                if not sheet_name or not cell_ref:
                    continue

                if sheet_name not in wb.sheetnames:
                    failed_cells.append(f"工作表「{sheet_name}」不存在")
                    continue

                ws = wb[sheet_name]
                try:
                    ws[cell_ref] = value
                    filled_count += 1
                except:
                    col_str = "".join(c for c in cell_ref if c.isalpha())
                    row_str = "".join(c for c in cell_ref if c.isdigit())
                    if col_str and row_str:
                        try:
                            col_idx = openpyxl.utils.column_index_from_string(col_str)
                            row_idx = int(row_str)
                            ws.cell(row=row_idx, column=col_idx, value=value)
                            filled_count += 1
                        except:
                            failed_cells.append(cell_ref)

        elif isinstance(data, dict):
            for sheet_name, cells in data.items():
                if sheet_name not in wb.sheetnames:
                    failed_cells.append(f"工作表「{sheet_name}」不存在")
                    continue
                ws = wb[sheet_name]
                if isinstance(cells, dict):
                    for cell_ref, value in cells.items():
                        if value is None:
                            continue
                        try:
                            ws[cell_ref] = value
                            filled_count += 1
                        except:
                            failed_cells.append(cell_ref)

        wb.save(str(out_path))
        wb.close()

        return json.dumps({
            "status": "filled",
            "output_path": str(out_path),
            "output_filename": out_name,
            "filled_count": filled_count,
            "failed_cells": failed_cells,
        }, ensure_ascii=False, default=str)

    except Exception as e:
        return json.dumps({"error": f"填充失败: {str(e)}"}, ensure_ascii=False)


@registry.register(name="batch_extract", category="file")
def batch_extract(folder_path: str, keywords: str = "") -> str:
    import openpyxl
    from backend.tools.excel_reader import read_excel
    from backend.tools.pdf_reader import read_pdf

    settings = get_settings()
    target = Path(folder_path)
    if not target.exists():
        target = Path(settings.data_dir) / folder_path
    if not target.exists():
        return json.dumps({"error": f"文件夹不存在: {folder_path}"}, ensure_ascii=False)

    keyword_list = [k.strip() for k in keywords.split(",") if k.strip()] if keywords else []

    all_data = {}
    file_list = sorted(target.rglob("*")) if target.is_dir() else [target]

    for f in file_list:
        if not f.is_file():
            continue
        ext = f.suffix.lower()
        file_key = str(f.relative_to(target)) if target.is_dir() else f.name

        if ext in {".xlsx", ".xls", ".csv"}:
            result = read_excel(str(f))
            try:
                all_data[file_key] = json.loads(result)
            except:
                all_data[file_key] = {"raw": result}
        elif ext == ".pdf":
            result = read_pdf(str(f))
            try:
                all_data[file_key] = json.loads(result)
            except:
                all_data[file_key] = {"raw": result}
        elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".tiff"}:
            all_data[file_key] = {"type": "image", "note": "需要OCR处理"}

    if keyword_list:
        filtered = {}
        for fname, fdata in all_data.items():
            data_str = json.dumps(fdata, ensure_ascii=False, default=str)
            if any(kw in data_str for kw in keyword_list):
                filtered[fname] = fdata
        all_data = filtered

    return json.dumps({
        "source_folder": str(target),
        "total_files_extracted": len(all_data),
        "data": all_data,
    }, ensure_ascii=False, default=str)
